import pypandoc
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os
import win32com.client

print("=" * 60)
print("BUILDING COMPLETE MSC DISSERTATION WITH FULL PREAMBLE")
print("=" * 60)

# Step 1: Combine Markdown files (Preamble + Chapters 1-6)
md_files = [
    '0_Preamble_Draft.md',
    'Chapter_1_Draft.md',
    'Chapter_2_Draft.md',
    'Chapter_3_Draft.md',
    'Chapter_4_Draft.md',
    'Chapter_5_Draft.md',
    'Chapter_6_Draft.md'
]

combined_md = ""
for file in md_files:
    print(f"Reading {file}...")
    with open(file, 'r', encoding='utf-8') as f:
        combined_md += f.read() + "\n\n"

with open('Full_Draft.md', 'w', encoding='utf-8') as f:
    f.write(combined_md)

# Step 2: Convert combined Markdown to Word using Pandoc with reference template
print("\nConverting Full_Draft.md to DOCX via Pandoc...")
temp_docx = 'temp_dissertation.docx'
pypandoc.convert_file(
    'Full_Draft.md', 
    'docx', 
    outputfile=temp_docx, 
    extra_args=['--reference-doc=report temp.docx', '--resource-path=visuals']
)

# Step 3: Polish and Format with python-docx
print("\nPolishing document with python-docx...")
doc = docx.Document(temp_docx)

# Style helper: shade table header cells
def set_header_cell_shading(cell, color_hex="1E293B"):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_row_shading(cell, color_hex="F8FAFC"):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

# Format all tables
print(f"Formatting {len(doc.tables)} tables in the document...")
for t_idx, table in enumerate(doc.tables):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Format Header row
    for cell in table.rows[0].cells:
        set_header_cell_shading(cell, "1E293B")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
                
    # Format data rows
    for r_idx in range(1, len(table.rows)):
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for cell in table.rows[r_idx].cells:
            if bg_color == "F8FAFC":
                set_row_shading(cell, bg_color)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)

# Step 4: Add Page Breaks Before Major Headings
major_headings = [
    "Table of Contents",
    "List of Figures",
    "List of Tables",
    "Abstract",
    "Declaration",
    "Acknowledgements",
    "Abbreviations",
    "Chapter 1:",
    "Chapter 2:",
    "Chapter 3:",
    "Chapter 4:",
    "Chapter 5:",
    "Chapter 6:",
    "References"
]

print("Applying page breaks before major sections...")
for p in doc.paragraphs:
    text = p.text.strip()
    for heading in major_headings:
        if (text == heading or text.startswith(heading)) and len(text) < 55:
            p.paragraph_format.page_break_before = True
            break

# Format Title Page (first few paragraphs)
print("Styling Title Page...")
for p in doc.paragraphs[:15]:
    if "Title Page" in p.text:
        p.text = "IMAGE CLASSIFICATION USING MACHINE LEARNING FOR DIGITAL FORENSIC INVESTIGATIONS"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(18)
            run.font.bold = True
    elif "Project Title:" in p.text:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif "Author:" in p.text or "Student ID:" in p.text or "Degree:" in p.text or "Supervisor:" in p.text:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Step 5: Save to MSc_Dissertation_Draft.docx
output_path = os.path.abspath('MSc_Dissertation_Draft.docx')
doc.save(output_path)
print(f"\nDocument saved successfully to {output_path}")

# Cleanup temp files
if os.path.exists('Full_Draft.md'): os.remove('Full_Draft.md')
if os.path.exists(temp_docx): os.remove(temp_docx)

print("\n" + "=" * 60)
print("BUILD COMPLETE! Preamble (TOC, Figures, Tables, Abbreviations) is fully integrated.")
print("=" * 60)
