import docx

doc = docx.Document('Pereira_Alston_24850898.docx')

def dump_range(start, end, filename):
    with open(filename, 'w', encoding='utf-8') as f:
        for i in range(start, min(end, len(doc.paragraphs))):
            p = doc.paragraphs[i]
            txt = p.text.strip()
            if not txt:
                continue
            f.write(f"=== P{i} [{p.style.name}] ===\n")
            has_math = bool(p._p.xpath('.//m:oMath | .//m:oMathPara'))
            has_draw = bool(p._p.xpath('.//w:drawing'))
            f.write(f"Tags: math={has_math}, drawing={has_draw}, runs={len(p.runs)}\n")
            for r_i, r in enumerate(p.runs):
                flags = []
                if r.bold: flags.append('BOLD')
                if r.italic: flags.append('ITALIC')
                f.write(f"  r{r_i} [{'|'.join(flags)}]: {repr(r.text)}\n")
            f.write(f"FULL TEXT:\n{p.text}\n\n")

dump_range(8, 42, 'block1_ch1.txt')
dump_range(42, 103, 'block2_ch2.txt')
dump_range(103, 146, 'block3_ch3.txt')
dump_range(146, 191, 'block4_ch4.txt')
dump_range(191, 273, 'block5_ch5.txt')
dump_range(273, 305, 'block6_ch6.txt')

print("Dumped all blocks successfully.")
