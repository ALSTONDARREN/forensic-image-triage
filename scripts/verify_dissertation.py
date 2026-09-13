import docx, re

doc = docx.Document('Pereira_Alston_24850898.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Check drawings
drawings = 0
for i, p in enumerate(doc.paragraphs):
    d = p._p.xpath('.//w:drawing')
    if d:
        drawings += len(d)
print(f"Total drawings in paragraphs: {drawings}")

# Check math
math_paras = 0
for i, p in enumerate(doc.paragraphs):
    m = p._p.xpath('.//m:oMath | .//m:oMathPara')
    if m:
        math_paras += 1
print(f"Total paragraphs with math: {math_paras}")

# Check remaining em/en dashes
dashes_found = []
for i, p in enumerate(doc.paragraphs):
    # skip references (P305+) where en-dash in book/conference titles is standard bibliography
    if i >= 305:
        continue
    if '—' in p.text or '–' in p.text:
        dashes_found.append((i, p.text[:100]))

print(f"Remaining dashes before references: {len(dashes_found)}")
for i, txt in dashes_found:
    print(f"  P{i}: {repr(txt)}")

# Check remaining non-heading bold runs
bolds_found = []
for i, p in enumerate(doc.paragraphs):
    if p.style.name.startswith('Heading') or p.style.name.startswith('Title'):
        continue
    if i == 0: # Title page labels
        continue
    bolds = [r.text for r in p.runs if r.bold and r.text.strip()]
    if bolds:
        bolds_found.append((i, bolds))

print(f"Remaining non-heading bold runs in body: {len(bolds_found)}")
for i, bolds in bolds_found:
    print(f"  P{i}: {bolds}")

# Check for any corrupted characters (\ufffd)
corrupted = []
for i, p in enumerate(doc.paragraphs):
    if '\ufffd' in p.text:
        corrupted.append(i)
print(f"Paragraphs with corrupted chars: {len(corrupted)}")
