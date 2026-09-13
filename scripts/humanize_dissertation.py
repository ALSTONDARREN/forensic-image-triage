import docx
import shutil

# Make sure we have a fresh copy from backup
shutil.copyfile('Pereira_Alston_24850898_BACKUP.docx', 'Pereira_Alston_24850898.docx')
doc = docx.Document('Pereira_Alston_24850898.docx')

print(f"Loaded Pereira_Alston_24850898.docx with {len(doc.paragraphs)} paragraphs.")

# Helper to remove bold from all runs in a paragraph
def unbold_paragraph(p):
    for r in p.runs:
        r.bold = False

# Helper to replace text in runs preserving italics/math
def replace_in_runs(p, search_str, replace_str):
    for r in p.runs:
        if search_str in r.text:
            r.text = r.text.replace(search_str, replace_str)

# Helper to set simple paragraph text without bold, optionally with selective italics
def set_para_clean_text(p, text, italic_phrases=None):
    p.text = "" # clears runs in this paragraph
    unbold_paragraph(p)
    if not italic_phrases:
        r = p.add_run(text)
        r.bold = False
    else:
        # split by italic phrases
        current_text = text
        while current_text:
            earliest_idx = len(current_text)
            chosen_phrase = None
            for phrase in italic_phrases:
                idx = current_text.find(phrase)
                if idx != -1 and idx < earliest_idx:
                    earliest_idx = idx
                    chosen_phrase = phrase
            if chosen_phrase is not None:
                if earliest_idx > 0:
                    r_pre = p.add_run(current_text[:earliest_idx])
                    r_pre.bold = False
                r_it = p.add_run(chosen_phrase)
                r_it.italic = True
                r_it.bold = False
                current_text = current_text[earliest_idx + len(chosen_phrase):]
            else:
                r_post = p.add_run(current_text)
                r_post.bold = False
                current_text = ""

# -------------------------------------------------------------
# P9: Abstract
# -------------------------------------------------------------
p9_text = (
    "Digital forensic investigators face growing challenges managing the expanding volume of "
    "digital evidence seized during criminal investigations. Although deep learning models provide "
    "high classification accuracy, their significant computational overhead and lack of algorithmic "
    "transparency present substantial barriers for frontline forensic units operating standard CPU-based "
    "hardware. This dissertation evaluates a hybrid evidence triage pipeline combining classical computer "
    "vision techniques with Optical Character Recognition (OCR). By extracting deterministic 4292-dimensional "
    "structural (Histogram of Oriented Gradients) and colour (HSV) descriptors, a classical Support Vector "
    "Machine (SVM) classifier achieved an accuracy of 88.4% with a mean CPU inference latency of approximately "
    "96.5 ms, representing a threefold speed improvement over a ResNet-50 baseline (294.2 ms). To address "
    "visual ambiguities inherent to classical feature representations, such as misclassifying pharmaceutical "
    "containers as alcohol bottles, the pipeline integrates an OCR verification layer to extract contextual "
    "textual indicators. The findings demonstrate that hybrid classical pipelines offer significant latency "
    "advantages alongside methodological explainability, providing a practical and auditable solution for "
    "resource-constrained digital forensic triage."
)
set_para_clean_text(doc.paragraphs[9], p9_text)

# -------------------------------------------------------------
# Chapter 1: Introduction
# -------------------------------------------------------------
# P23
p23_text = (
    "Digital forensic investigators face increasing challenges arising from the rapid growth of seized "
    "digital media. With the widespread adoption of multi-terabyte storage drives, high-capacity smartphones, "
    "and automated cloud backups, routine criminal investigations frequently yield hundreds of gigabytes or "
    "terabytes of uncurated files. Locating specific visual evidence, such as images of illicit narcotics, "
    "prohibited weapons, or fraudulent documentation, presents a severe operational bottleneck when relying "
    "primarily on manual review."
)
set_para_clean_text(doc.paragraphs[23], p23_text)

# P24
p24_text = (
    "The theoretical foundation of evidence recovery is grounded in Locard’s Exchange Principle. Edmond Locard "
    "postulated the fundamental maxim: “Every contact leaves a trace” (Locard, 1930). In physical forensic science, "
    "this manifests as trace DNA, textile fibres, or latent fingerprints deposited at a scene. In digital forensics, "
    "Locard’s principle applies equally to digital artifacts. Whenever an individual interacts with a computer system, "
    "downloads an illicit file, or transmits data across a network, system activity records corresponding artifacts "
    "within file system metadata, the Windows Registry, or unallocated clusters. The primary operational challenge "
    "facing modern investigators is rarely the absence of artifacts, but rather isolating relevant evidentiary items "
    "from the extensive background data generated by standard operating system processes."
)
set_para_clean_text(doc.paragraphs[24], p24_text, italic_phrases=["“Every contact leaves a trace”"])

# P25
p25_text = (
    "Historically, forensic investigations have relied on cryptographic hash matching alongside manual examination "
    "to identify known illicit files. Algorithms such as Message Digest 5 (MD5) and Secure Hash Algorithm 1 (SHA-1) "
    "generate fixed-length digests representing digital fingerprints of target files. These digests are cross-referenced "
    "against established repositories of known material, such as the National Software Reference Library (NSRL). "
    "Although computationally efficient, exact cryptographic matching is vulnerable to minor file modifications due "
    "to the avalanche effect. If an image is cropped by a single pixel, re-encoded under lossy compression, or stripped "
    "of its EXIF metadata, the resulting digest changes entirely, preventing automated identification."
)
set_para_clean_text(doc.paragraphs[25], p25_text)

# P26
p26_text = (
    "Furthermore, the security assurances of legacy hashing algorithms have weakened significantly. MD5 has been "
    "susceptible to practical collision attacks since 2004, and the 2017 SHAttered demonstration by Stevens et al. (2017) "
    "established that distinct files could generate identical SHA-1 digests. When suspects employ anti-forensic "
    "adjustments to alter file signatures, automated lookup tables fail, necessitating manual review. Manual verification "
    "is exceptionally time-consuming, resource-intensive, and subjects forensic analysts to prolonged cognitive fatigue "
    "and potential psychological strain when examining disturbing or illicit material."
)
set_para_clean_text(doc.paragraphs[26], p26_text)

# P27
p27_text = (
    "In recent years, deep learning approaches, particularly Convolutional Neural Networks (CNNs), have been widely "
    "investigated as potential solutions to forensic backlogs. While CNNs provide high feature abstraction capabilities, "
    "they present notable operational constraints for frontline law enforcement. Deep neural architectures require "
    "extensive labelled datasets, exhibit limited interpretability due to complex non-linear parameter spaces, and "
    "typically rely on specialized Graphical Processing Units (GPUs) for timely inference, resources that are rarely "
    "standard across regional forensic units."
)
set_para_clean_text(doc.paragraphs[27], p27_text)

# P28
p28_text = (
    "The practical impact of evidence backlogs is well documented in criminal case reviews. In the 2017 case of "
    "R v Allan (Crown Prosecution Service, 2018), a criminal trial in the United Kingdom collapsed after investigators "
    "failed to process and disclose a relevant cache of approximately 40,000 phone messages prior to proceedings. "
    "The delayed identification of exculpatory material led to the abandonment of the prosecution and drew scrutiny "
    "to digital evidence handling procedures. When forensic units are unable to process incoming digital media within "
    "operational timeframes, trials face protracted delays. This highlights the practical necessity for rapid, "
    "automated, and verifiable evidence triage tools that assist human investigators early in the investigative lifecycle."
)
set_para_clean_text(doc.paragraphs[28], p28_text, italic_phrases=["R v Allan"])

# P29
p29_text = (
    "Classical computer vision techniques based on deterministic feature descriptors provide an efficient, explainable "
    "alternative for early-stage triage. By calculating structured spatial gradients, such as the Histogram of Oriented "
    "Gradients (HOG; Dalal and Triggs, 2005), alongside global colour distributions, systems can classify visual evidence "
    "on standard CPU workstations without dedicated hardware acceleration. Nonetheless, purely visual classifiers "
    "remain vulnerable to systematic ambiguities, such as misclassifying a brown glass medicine bottle as an alcoholic "
    "beverage due to shared geometry and colour. To address these edge cases, forensic triage can be structured as a "
    "hybrid framework, combining the computational efficiency of classical descriptor extraction with the contextual "
    "validation of Optical Character Recognition (OCR)."
)
set_para_clean_text(doc.paragraphs[29], p29_text)

# P34, P35, P36
p34_text = (
    "This research is motivated by the need to improve the efficiency of forensic image triage while avoiding the high "
    "computational requirements and opacity associated with unconstrained deep learning models."
)
set_para_clean_text(doc.paragraphs[34], p34_text)

p35_text = (
    "Digital forensics and incident response units frequently operate under stringent investigative deadlines and finite "
    "budgets. Investigators cannot feasibly spend weeks conducting exhaustive manual reviews of hundreds of thousands of "
    "benign media files, such as system graphics or personal photographs, to isolate relevant evidentiary artifacts. "
    "At the same time, regional police constabularies rarely possess dedicated GPU computing clusters to deploy large "
    "deep learning pipelines across all field units."
)
set_para_clean_text(doc.paragraphs[35], p35_text)

p36_text = (
    "By developing a structured triage pipeline based on classical computer vision, this project aims to accelerate the "
    "initial screening phase of digital investigations. A lightweight, deterministic classification tool executing "
    "locally on standard forensic workstations can filter non-evidentiary media, allowing examiners to focus investigative "
    "attention on flagged items while reducing manual exposure to potentially distressing content."
)
set_para_clean_text(doc.paragraphs[36], p36_text)

# P39: Objectives - remove unwanted bolding, format cleanly
p39_text = (
    "To achieve this aim, the research is guided by the following specific objectives: "
    "1. Dataset Engineering: Curate and normalize a representative dataset of digital images across three primary "
    "forensic categories (Alcohol, Drugs, and Weapons), incorporating varied lighting, occlusions, and backgrounds "
    "to reflect authentic investigative conditions. "
    "2. Classical Feature Extraction: Develop an efficient preprocessing and feature extraction pipeline using OpenCV "
    "to extract deterministic 3D HSV colour histograms and Histogram of Oriented Gradients (HOG) structural descriptors. "
    "3. Machine Learning Implementation: Train and evaluate three classical classification algorithms (Support Vector "
    "Machine, Random Forest, and K-Nearest Neighbors) using the engineered 4292-dimensional feature space. "
    "4. Baseline Comparison: Benchmark the classical OpenCV pipeline against three transfer-learned Convolutional "
    "Neural Networks (ResNet-50, VGG-16, and MobileNetV3-Large), comparing classification metrics directly against CPU "
    "inference latency. "
    "5. Hybrid Semantic Integration: Incorporate an OCR validation mechanism using EasyOCR to identify textual markings "
    "(such as pharmaceutical labels and weapon designations) to mitigate visual misclassification in ambiguous cases. "
    "6. Deployment: Integrate the classification and OCR modules into a functional triage interface capable of executing "
    "locally on standard forensic workstations."
)
set_para_clean_text(doc.paragraphs[39], p39_text)

# P41: Structure - remove unwanted bolding, remove stray markdown asterisks
p41_text = (
    "To present the investigation systematically, the remainder of this dissertation is organized as follows: "
    "Chapter 2 (Literature Review) critically reviews automated digital forensic triage, contrasting the limitations of "
    "cryptographic hashing with the computational and legal challenges of deep learning, and establishing the theoretical "
    "basis for deterministic computer vision. "
    "Chapter 3 (Design) details the experimental design, dataset curation, and mathematical formulation of the HOG and HSV "
    "feature descriptors alongside the deep learning baselines. "
    "Chapter 4 (Implementation) outlines the technical implementation of the OpenCV classification pipeline, the transfer "
    "learning configuration, and the integration of the OCR verification layer. "
    "Chapter 5 (Evaluation) provides an empirical analysis of classification accuracy, ROC/AUC metrics, granular class "
    "confusion, and CPU inference latency across the evaluated models. "
    "Chapter 6 (Conclusion) synthesizes the research findings, discusses operational and socioeconomic implications, addresses "
    "limitations, and identifies directions for future work."
)
set_para_clean_text(doc.paragraphs[41], p41_text)

# -------------------------------------------------------------
# Chapter 2: Literature Review
# -------------------------------------------------------------
# P44
p44_text = (
    "Digital forensics has changed significantly over the past two decades as digital devices have become central to criminal "
    "investigations. The expansion of consumer storage and communication platforms has led to substantial growth in the volume "
    "of data seized during investigations. Consequently, law enforcement agencies require automated, scalable triage "
    "methodologies to rapidly locate relevant evidentiary material, such as illicit narcotics, prohibited firearms, and "
    "contraband imagery. This chapter reviews the academic literature and operational standards surrounding automated evidence "
    "triage. It examines the conventional use of cryptographic hash matching, evaluates the operational and legal constraints "
    "associated with deep neural networks, and establishes the theoretical justification for employing deterministic classical "
    "computer vision supplemented by semantic text verification."
)
set_para_clean_text(doc.paragraphs[44], p44_text)

# P47
p47_text = (
    "While hash matching is computationally efficient and universally accepted in legal proceedings as a method for verifying "
    "evidential integrity (chain of custody), it exhibits key limitations when deployed as an investigative search tool due to the "
    "avalanche effect. Cryptographic hashes are designed such that the slightest variation in input data yields an entirely distinct "
    "digest. If a file is cropped by a single pixel, compressed, or stripped of its EXIF metadata during transmission, the resulting "
    "hash value changes completely, preventing automated database matching. Because suspects frequently manipulate digital media to "
    "evade signature detection, exact hash matching alone cannot provide a comprehensive triage solution."
)
set_para_clean_text(doc.paragraphs[47], p47_text)

# P48
p48_text = (
    "Furthermore, the security assurances of legacy hashing algorithms have weakened significantly. MD5 has been vulnerable to "
    "practical collision attacks since 2004, enabling the generation of distinct files producing identical hash values. Similarly, "
    "SHA-1 was demonstrated to be vulnerable in 2017 with the publication of the SHAttered attack by Stevens et al. (2017). "
    "The existence of collision vulnerabilities introduces evidential complications; if defense counsel demonstrates that an "
    "algorithm is no longer strictly collision-resistant, the foundational integrity of the digital evidence may be questioned. "
    "Due to these sensitivity and collision limitations, exact hashing cannot serve as a comprehensive search mechanism for altered "
    "files, requiring investigators to conduct manual examinations. Manual inspection is resource-intensive and exposes examiners "
    "to cognitive fatigue and psychological strain when reviewing explicit or distressing content over extended periods."
)
set_para_clean_text(doc.paragraphs[48], p48_text)

# P52
p52_text = (
    "However, an evidential paradox arises when investigators attempt to utilize cryptographic hashing as an active search "
    "mechanism (such as querying seized files against a database of known contraband). While a hash reliably confirms that a file "
    "has not been modified by the police, it cannot detect illicit files that have undergone minor adjustments by the suspect. "
    "If an offender acquires a known illicit image and slightly modifies contrast or resizes the canvas, the visual content remains "
    "contraband, but the cryptographic hash changes entirely, rendering database queries ineffective. The inability of hash functions "
    "to capture semantic similarity rather than exact bit-level duplication is a primary driver behind the adoption of computer "
    "vision in forensic triage."
)
set_para_clean_text(doc.paragraphs[52], p52_text, italic_phrases=["police", "suspect"])

# P55: Notice P55 has math tags (<m:oMath> for 3x3). Let's edit the text runs directly!
p55 = doc.paragraphs[55]
unbold_paragraph(p55)
p55.runs[0].text = (
    "The modern adoption of deep learning in computer vision gained widespread momentum in 2012 following the introduction of "
    "AlexNet (Krizhevsky et al., 2012), which achieved competitive error rates on the ImageNet challenge using deep convolutional "
    "layers and Rectified Linear Unit (ReLU) activations. This was followed by deeper architectures such as VGG-16 (Simonyan and "
    "Zisserman, 2014), which utilized uniform "
)
p55.runs[1].text = (
    " convolution filters to extract dense hierarchical feature maps. Subsequently, the vanishing gradient problem (which historically "
    "constrained the depth of trainable feedforward networks) was mitigated by the development of Residual Networks (ResNet) in 2015 "
    "(He et al., 2016). ResNet incorporated identity skip connections that permit gradients to propagate directly across layers during "
    "backpropagation, facilitating the stable training of networks with dozens or hundreds of layers."
)

# P56
p56_text = (
    "Recent empirical research into automated evidence processing highlights the utility of machine learning approaches; for instance, "
    "Del Mar-Raave et al. (2021) developed a machine learning-based forensic tool for image classification adopting a design science "
    "approach, demonstrating that automated visual categorization can substantially reduce investigator backlogs compared to traditional "
    "manual review. Academic studies have consistently shown that models such as ResNet-50 and VGG-16 achieve high accuracy when "
    "classifying specialized evidentiary objects, including firearms and illicit pharmaceutical packaging. However, deploying these "
    "architectures within operational digital forensic units introduces practical challenges regarding hardware availability, inference "
    "throughput, and legal explainability."
)
set_para_clean_text(doc.paragraphs[56], p56_text)

# P59
p59_text = (
    "Deep learning architectures are notoriously data-intensive, requiring extensive, diverse, and accurately annotated datasets to "
    "generalize effectively. Within digital forensics, acquiring large, diverse, and representative corpora of authentic illicit media is "
    "subject to strict legal, regulatory, and ethical constraints. Privacy frameworks such as the General Data Protection Regulation "
    "(GDPR), judicial disclosure restrictions, and the sensitive nature of evidential material (such as Child Sexual Abuse Material or "
    "confidential corporate records) prohibit public dissemination and open sharing of real-world forensic datasets."
)
set_para_clean_text(doc.paragraphs[59], p59_text)

# P60
p60_text = (
    "Consequently, forensic deep learning models are frequently trained on constrained, imbalanced, or synthetic datasets. When deployed "
    "operationally, these models risk overfitting and may underperform when presented with novel variations in evidence that diverge "
    "from their training distributions."
)
set_para_clean_text(doc.paragraphs[60], p60_text)

# P62
p62_text = (
    "Data scarcity can also introduce algorithmic bias into trained models, presenting ethical and legal concerns in criminal proceedings. "
    "If a deep neural network is trained disproportionately on contraband images captured within particular domestic environments, the "
    "model may inadvertently learn spurious correlations between background features (such as interior furnishings, wall textures, or "
    "ambient lighting) and illicit classifications. When deployed, the system risks generating skewed false-positive rates based on ambient "
    "scenery rather than the physical attributes of the contraband itself."
)
set_para_clean_text(doc.paragraphs[62], p62_text)

# P64
p64_text = (
    "While object classification (such as identifying a handgun) is ostensibly less susceptible to demographic bias than facial "
    "recognition, the underlying risk of environmental bias remains notable. Because deep CNNs learn abstract, holistic representations "
    "across the full image, they remain vulnerable to spurious background correlations. In contrast, classical models using the Histogram "
    "of Oriented Gradients (HOG) mitigate this risk by design. Because HOG explicitly models the local gradient orientation and edge "
    "structure of the foreground object (such as the silhouette of a firearm or container) while discarding diffuse background color "
    "fields, it exhibits lower susceptibility to learning non-causal environmental correlations. Consequently, deterministic classical "
    "feature representations provide both computational predictability and a transparent foundation for evidential triage in criminal "
    "justice contexts."
)
set_para_clean_text(doc.paragraphs[64], p64_text, italic_phrases=["foreground object"])

# P66
p66_text = (
    "A major practical barrier to adopting deep neural networks in operational law enforcement is the computational requirement for "
    "model execution. Training deep CNNs and processing large batches of seized media typically demands dedicated Graphics Processing "
    "Units (GPUs)."
)
set_para_clean_text(doc.paragraphs[66], p66_text)

# P67
p67_text = (
    "Many deep learning models are characterized by high parameter counts that require substantial memory and matrix compute. In practice, "
    "frontline digital forensic units and field investigators frequently operate on standard, CPU-equipped workstations without high-performance "
    "GPU accelerators. When large CNN architectures execute inference exclusively on CPU hardware, per-image latency increases substantially, "
    "turning rapid media screening into a protracted bottleneck."
)
set_para_clean_text(doc.paragraphs[67], p67_text)

# P69
p69_text = (
    "Furthermore, deep learning introduces significant challenges regarding algorithmic interpretability. A primary requirement for "
    "employing automated software outputs in legal proceedings (informed by the Daubert standard in the United States and equivalent "
    "reliability guidelines in the UK) is that evidential reasoning must be transparent, auditable, and scientifically grounded."
)
set_para_clean_text(doc.paragraphs[69], p69_text)

# P70
p70_text = (
    "Deep neural networks function largely as black-box systems. Although they produce definitive classification probabilities, the "
    "underlying interactions across millions of weighted parameters make it challenging for a forensic examiner to explain transparently "
    "how a particular prediction was generated. If a model incorrectly flags a benign photograph as contraband, an examiner cannot easily "
    "trace the neural pathway responsible. In judicial contexts, opaque automated outputs are vulnerable to evidential challenge by expert "
    "witnesses, potentially compromising investigative credibility."
)
set_para_clean_text(doc.paragraphs[70], p70_text, italic_phrases=["how"])

# P72
p72_text = (
    "In addition to interpretability challenges, deep neural architectures present potential vulnerabilities to adversarial manipulation. "
    "Because CNNs rely on high-dimensional non-linear feature maps, they can be misled by engineered perturbations that remain imperceptible "
    "to the human eye."
)
set_para_clean_text(doc.paragraphs[72], p72_text)

# P77
p77_text = (
    "In a digital forensic setting, adversarial techniques could theoretically be applied to evade automated detection. For example, "
    "small perturbations designed to alter high-level CNN activations could cause an automated classifier to misidentify a weapon or "
    "narcotic as a benign object. Because classical descriptors like HOG rely on aggregate gradient histograms over localized pixel blocks "
    "rather than deep non-linear feature maps, they are less vulnerable to imperceptible, single-pixel gradient shifts, providing an "
    "auditable and robust baseline for evidential screening."
)
set_para_clean_text(doc.paragraphs[77], p77_text)

# P79
p79_text = (
    "Issues of model opacity and vulnerability directly intersect with standards for evidential admissibility. Under the Frye standard "
    "(Frye v. United States, 1923), scientific methods must have achieved general acceptance within the relevant field. Under the Daubert "
    "standard (Daubert v. Merrell Dow Pharmaceuticals, Inc., 1993), which informs international evidence principles, admissibility assesses "
    "whether a technique has been empirically tested, subjected to peer review, possesses an established error rate, and adheres to "
    "operational standards."
)
set_para_clean_text(doc.paragraphs[79], p79_text, italic_phrases=["Frye v. United States", "Daubert v. Merrell Dow Pharmaceuticals, Inc."])

# P80
p80_text = (
    "When automated classifications are introduced in court, opaque neural network decisions can present evidential challenges. If an "
    "investigator cannot explain how a specific classification was derived, or demonstrate the exact boundaries of the decision space, "
    "the reliability of the evidence may be contested. In contrast, classical computer vision pipelines rely on explicit mathematical "
    "descriptors. An expert witness can demonstrate that an SVM identified an object because the feature extractor recorded defined "
    "gradient orientations corresponding to the physical geometry of the item, providing a clear, verifiable chain of reasoning."
)
set_para_clean_text(doc.paragraphs[80], p80_text, italic_phrases=["how"])

# P82
p82_text = (
    "Given the hardware demands and interpretability constraints of complex neural networks, classical computer vision techniques "
    "implemented in optimized libraries such as OpenCV offer a practical alternative for digital forensic triage."
)
set_para_clean_text(doc.paragraphs[82], p82_text)

# P85
p85_text = (
    "Early classical computer vision relied heavily on keypoint detectors such as the Scale-Invariant Feature Transform (SIFT; Lowe, 2004) "
    "and Speeded-Up Robust Features (SURF; Bay et al., 2006). SIFT introduced robust detection of distinctive local interest points "
    "invariant to scale, rotation, and illumination changes (Lowe, 2004). However, SIFT was computationally intensive and restricted by "
    "patents for many years. SURF improved throughput via Haar wavelet approximations (Bay et al., 2006), but still imposed notable "
    "computational overhead for high-throughput triage."
)
set_para_clean_text(doc.paragraphs[85], p85_text)

# P87
p87_text = (
    "While local interest point detectors excel at identifying localized corners, capturing the structural geometry of objects such "
    "as firearms or drug paraphernalia requires dense contour analysis. The Histogram of Oriented Gradients (HOG) method, developed "
    "initially for pedestrian detection (Dalal and Triggs, 2005), evaluates the distribution of local intensity gradient directions across "
    "an image."
)
set_para_clean_text(doc.paragraphs[87], p87_text)

# P88
p88_text = (
    "By subdividing an image into small connected regions (cells) and compiling histograms of gradient orientations for pixels within "
    "each cell, HOG maps the structural boundary of an object independently of minor color variations (Dalal and Triggs, 2005). HOG is "
    "unencumbered by proprietary licensing and computationally efficient on standard CPU hardware, making it well suited for structured "
    "shape detection in forensic triage applications."
)
set_para_clean_text(doc.paragraphs[88], p88_text)

# P90
p90_text = (
    "While structural geometry is critical, contraband items frequently display recognizable colour characteristics. For instance, "
    "pharmaceutical preparations and commercial packaging rely on standardized colour palettes. Standard RGB representations are "
    "sensitive to lighting variations, such as shadows or overexposure. In contrast, the Hue, Saturation, Value (HSV) colour space "
    "decouples chromaticity (Hue) from illumination intensity (Value). By calculating multi-dimensional HSV colour histograms, "
    "forensic triage systems can incorporate lighting-invariant colour distributions with minimal computational overhead."
)
set_para_clean_text(doc.paragraphs[90], p90_text)

# P92
p92_text = (
    "A documented limitation of appearance-based classification is the potential for visual ambiguity among structurally similar "
    "objects. For example, a classical model evaluated on brown cylindrical glass containers may classify other cylindrical items, "
    "such as prescription medicine bottles or wooden handles, as alcoholic beverages due to shared geometry and colour."
)
set_para_clean_text(doc.paragraphs[92], p92_text)

# P93
p93_text = (
    "To resolve visual ambiguities, modern forensic research emphasizes hybrid methodologies incorporating textual context. Images "
    "containing contraband often display key contextual indicators, such as pharmaceutical markings, chemical terminology, or manufacturer "
    "markings on weapons, which purely geometric models cannot interpret directly."
)
set_para_clean_text(doc.paragraphs[93], p93_text)

# P94
p94_text = (
    "Optical Character Recognition (OCR), utilizing modern architectures such as Character Region Awareness for Text Detection (CRAFT; "
    "Baek et al., 2019) and Convolutional Recurrent Neural Networks (CRNN; Shi et al., 2016), provides a valuable semantic verification "
    "layer within forensic pipelines. By executing OCR alongside primary visual classification, a triage framework can extract contextual "
    "text from object surfaces. If an image of pharmaceutical syrup is initially categorized as alcohol based on its container geometry, "
    "the OCR module can identify terms such as 'prescription' or 'codeine' and programmatically update the classification to the narcotics "
    "category."
)
set_para_clean_text(doc.paragraphs[94], p94_text)

# P99
p99_text = (
    "As summarized in Table 2.1, existing methodologies involve distinct operational tradeoffs: cryptographic hashing provides microsecond "
    "lookup speeds but fails when files undergo minor alterations, whereas deep CNNs provide high visual invariance at the cost of substantial "
    "CPU inference latency and reduced interpretability. The hybrid framework evaluated in this dissertation seeks to balance these "
    "considerations by combining lightweight HOG/HSV computer vision descriptors with an explicit OCR verification mechanism."
)
set_para_clean_text(doc.paragraphs[99], p99_text)

# P101: Summary - remove unwanted bold runs
p101_text = (
    "A review of the literature highlights the key operational tradeoffs in automated forensic triage: "
    "1. Cryptographic hashing (MD5/SHA-1) is computationally efficient but sensitive to minor bit-level alterations, limiting its utility "
    "for detecting modified visual evidence. "
    "2. Deep learning architectures provide strong feature abstraction across diverse imagery but incur substantial latency penalties "
    "on CPU hardware and present interpretability challenges under legal evidentiary standards. "
    "3. Classical computer vision utilizing HOG and HSV descriptors offers an explainable, computationally lightweight alternative "
    "capable of running on standard CPU hardware. "
    "Integrating deterministic visual descriptors with OCR text extraction provides a balanced approach, combining processing "
    "efficiency with semantic context to support frontline forensic triage."
)
set_para_clean_text(doc.paragraphs[101], p101_text)

# P102
p102_text = (
    "To address evidence backlogs in operational digital forensic laboratories, combining the CPU throughput of deterministic classical "
    "algorithms with the contextual verification of Optical Character Recognition offers a balanced, explainable, and practical approach "
    "to forensic image triage."
)
set_para_clean_text(doc.paragraphs[102], p102_text)

# -------------------------------------------------------------
# Chapter 3: Design
# -------------------------------------------------------------
# P105
p105_text = (
    "The primary objective of this research was to evaluate whether deterministic classical computer vision techniques retain practical "
    "value within resource-constrained digital forensic environments, or whether deep learning architectures represent the only viable "
    "solution. To assess this question empirically, an end-to-end forensic image triage framework was designed and implemented."
)
set_para_clean_text(doc.paragraphs[105], p105_text)

# P106
p106_text = (
    "The methodology directly addresses two operational constraints identified in the literature: the sensitivity of exact hash matching "
    "to minor file modifications, and the high computational overhead of deep neural networks on standard CPU workstations typically "
    "deployed in the field. Accordingly, the experimental design compares an OpenCV-based feature extraction pipeline against three "
    "benchmark Convolutional Neural Network (CNN) architectures evaluated under identical CPU execution constraints."
)
set_para_clean_text(doc.paragraphs[106], p106_text)

# P111
p111_text = (
    "To ensure all evaluated models were trained on representative data, the study required a corpus reflecting the unstructured "
    "characteristics of evidence seized in operational contexts. Standard academic benchmark datasets often feature objects isolated "
    "against clean backgrounds with controlled lighting. Training a forensic triage model exclusively on pristine or artificial datasets "
    "risks reduced generalization when evaluating unstructured, poorly illuminated, or partially occluded media seized during real-world "
    "investigations."
)
set_para_clean_text(doc.paragraphs[111], p111_text)

# P112: Categories - remove unwanted bolding
p112_text = (
    "The dataset focuses on three primary categories of forensic evidence: "
    "1. Alcohol: Commercial beverage containers, spirits bottles, and associated branded packaging. "
    "2. Drugs: Illicit narcotics, prescription medications, packaging, blister packs, and related paraphernalia. "
    "3. Weapons: Firearms (handguns, shotguns, rifles), bladed articles (knives, machetes), and associated offensive implements."
)
set_para_clean_text(doc.paragraphs[112], p112_text)

# P114
p114_text = (
    "Rather than relying strictly on pre-compiled academic collections, an automated scraping pipeline was implemented using Python. "
    "Utilizing web request libraries and search interfaces, image samples were gathered across varied online repositories. This collected "
    "imagery contained representative visual variations, including complex background textures, non-uniform lighting conditions, acute "
    "camera angles, and partial object occlusions."
)
set_para_clean_text(doc.paragraphs[114], p114_text)

# P123
p123_text = (
    "A key phase in developing the classical pipeline was the design of the explicit feature extraction process. Unlike deep learning "
    "architectures that derive latent representations during gradient-based training, classical machine learning relies on predefined "
    "mathematical descriptors that quantify visual attributes. An OpenCV-based extraction pipeline was implemented to ensure computational "
    "efficiency and deterministic feature generation."
)
set_para_clean_text(doc.paragraphs[123], p123_text)

# P129
p129_text = (
    "While colour distribution provides useful discriminatory information, structural geometry is essential for identifying physical "
    "implements. For example, firearms and bladed articles are characterized primarily by rigid contours and structural boundaries rather "
    "than surface pigmentation. The Histogram of Oriented Gradients (HOG) descriptor (Dalal and Triggs, 2005) was chosen over sparse "
    "keypoint detectors such as SIFT (Lowe, 2004) or SURF (Bay et al., 2006). Although SIFT provides invariance across isolated feature "
    "points, it incurs higher computational overhead. HOG captures the overall structural profile of an object across regular spatial "
    "grids, making it efficient for CPU-based classification."
)
set_para_clean_text(doc.paragraphs[129], p129_text)

# P136: Has math tags! Unbold run 5 ('3780 dimensions')
p136 = doc.paragraphs[136]
for r in p136.runs:
    r.bold = False

# P137: Unbold run 1 ('4292-dimensional')
p137 = doc.paragraphs[137]
for r in p137.runs:
    r.bold = False

# P145: Has math tags! Unbold run 3 ('4292-dimensional') and remove em-dashes
p145 = doc.paragraphs[145]
for r in p145.runs:
    r.bold = False
# Replace em-dashes in p145.runs[1]
p145.runs[1].text = p145.runs[1].text.replace(
    "—relying solely on inner products between support vectors—the full",
    " (relying solely on inner products between support vectors), the full"
)

# -------------------------------------------------------------
# Chapter 4: Implementation
# -------------------------------------------------------------
# P153
p153_text = (
    "With the 4292-dimensional feature space established, three deterministic machine learning algorithms were implemented using the "
    "scikit-learn library (Pedregosa et al., 2011) to serve as the classification models within the classical pipeline:"
)
set_para_clean_text(doc.paragraphs[153], p153_text)

# P154: Unbold run 0 ('Support Vector Machine (SVM):')
p154 = doc.paragraphs[154]
for r in p154.runs:
    r.bold = False

# P159: Has math tags! Polish text
p159 = doc.paragraphs[159]
p159.runs[3].text = p159.runs[3].text.replace(
    "Consequently, the SVM effectively ignores 95% of the background training data, constructing the decision boundary exclusively from the most critical, borderline forensic images.",
    "Consequently, the decision function depends exclusively on these margin-defining support vectors, contributing to inference efficiency."
)

# P160: Unbold run 0 ('Random Forest (RF):')
p160 = doc.paragraphs[160]
for r in p160.runs:
    r.bold = False

# P167: Remove em-dash
p167_text = (
    "By evaluating Information Gain across node splits, the Random Forest ensemble identifies informative feature dimensions, such as "
    "distinct hue ranges associated with specific pharmaceuticals, placing these splits near the root of individual decision trees."
)
set_para_clean_text(doc.paragraphs[167], p167_text)

# P168: Has math tags! Unbold run 0 ('K-Nearest Neighbors (KNN):')
p168 = doc.paragraphs[168]
for r in p168.runs:
    r.bold = False

# P171, P172, P173: Unbold deep learning model names
p171 = doc.paragraphs[171]
for r in p171.runs:
    r.bold = False

p172 = doc.paragraphs[172]
for r in p172.runs:
    r.bold = False

p173 = doc.paragraphs[173]
for r in p173.runs:
    r.bold = False

# P186: Has math tags! Unbold Accuracy, Precision, Recall, F1 Score; humanize run 9
p186 = doc.paragraphs[186]
for r in p186.runs:
    r.bold = False
p186.runs[9].text = p186.runs[9].text.replace(
    "In a forensic context, maximizing Recall is hyper-critical, as false negatives (missing a critical piece of illegal evidence) are infinitely more detrimental to an investigation than false positives.",
    "In a digital forensic context, recall is of particular operational importance, as false negatives (failing to detect contraband material) present greater investigative risk than false positives."
)

# P187: Unbold run 1 ('Inference Latency Benchmark') and remove em-dash in run 2
p187 = doc.paragraphs[187]
for r in p187.runs:
    r.bold = False
p187.runs[2].text = p187.runs[2].text.replace(
    "all models—both classical and PyTorch—to execute",
    "all models (both classical and PyTorch) to execute"
)

# P189: Remove em-dash
p189_text = (
    "A fundamental characteristic of closed-set classification is that models trained strictly on predefined evidentiary categories "
    "(alcohol, drugs, weapons) distribute class probabilities across those established classes via the softmax function. Consequently, "
    "when presented with an out-of-distribution (OOD) image, such as an everyday household item or outdoor scenery, closed-set classifiers "
    "will inevitably assign the item to one of the predefined classes based on superficial visual similarities."
)
set_para_clean_text(doc.paragraphs[189], p189_text)

# P190: Unbold runs 1 & 3, remove em-dash in run 2
p190 = doc.paragraphs[190]
for r in p190.runs:
    r.bold = False
p190.runs[2].text = p190.runs[2].text.replace(
    "rigid 50% threshold—indicating statistical uncertainty—the dashboard forcefully overrides the prediction and mathematically drops the image from the triage queue, classifying it as “Excluded (Out of Distribution)”.",
    "rigid 50% threshold (indicating statistical uncertainty), the dashboard filters out the image from the triage queue, categorizing it as “Excluded (Out of Distribution)”."
)

# -------------------------------------------------------------
# Chapter 5: Evaluation
# -------------------------------------------------------------
# P194
p194_text = (
    "This chapter presents and evaluates the empirical findings of the study. It first details the statistical classification "
    "performance (Accuracy, Precision, Recall, and F1-scores) of the classical OpenCV models (SVM, Random Forest, KNN) in comparison "
    "with the deep learning baselines (ResNet-50, VGG-16, MobileNetV3). Subsequently, it provides a granular, class-by-class error analysis, "
    "examining the visual characteristics that lead to misclassification. Finally, the chapter evaluates the CPU inference latency "
    "benchmarks, quantitatively assessing the throughput-to-accuracy tradeoff governing operational forensic deployment."
)
set_para_clean_text(doc.paragraphs[194], p194_text)

# P198: Unbold run 1 ('88.4%')
p198 = doc.paragraphs[198]
for r in p198.runs:
    r.bold = False

# P208: Unbold run 1 ('86.7%')
p208 = doc.paragraphs[208]
for r in p208.runs:
    r.bold = False

# P211: Has math! Unbold run 2 ('81.2%')
p211 = doc.paragraphs[211]
for r in p211.runs:
    r.bold = False

# P216: Unbold run 1 ('0.912')
p216 = doc.paragraphs[216]
for r in p216.runs:
    r.bold = False

# P228: Unbold run 1 ('99.50%')
p228 = doc.paragraphs[228]
for r in p228.runs:
    r.bold = False

# P232: Unbold run 1 ('99.50%')
p232 = doc.paragraphs[232]
for r in p232.runs:
    r.bold = False

# P236: Unbold run 1 ('99.60%')
p236 = doc.paragraphs[236]
for r in p236.runs:
    r.bold = False

# P249: Latencies list - unbold all runs
p249 = doc.paragraphs[249]
for r in p249.runs:
    r.bold = False
# Clean up markdown bullet asterisk in runs
for r in p249.runs:
    if '*' in r.text:
        r.text = r.text.replace('* ', '- ')

# P250: Extrapolations - unbold all runs
p250 = doc.paragraphs[250]
for r in p250.runs:
    r.bold = False

# P266, P267, P268: Case study - unbold headings, polish text
p266 = doc.paragraphs[266]
for r in p266.runs:
    r.bold = False
p266.runs[0].text = "Visual Misclassification:"

p267 = doc.paragraphs[267]
for r in p267.runs:
    r.bold = False
p267.runs[0].text = "Semantic Text Extraction:"

p268 = doc.paragraphs[268]
for r in p268.runs:
    r.bold = False
p268.runs[0].text = "Programmatic Correction:"
p268.runs[2].text = "revoked the SVM’s “Alcohol” prediction and re-categorized the evidence as “Drugs”."

# P269: Remove em-dash
p269_text = (
    "This case study demonstrates the utility of a multi-modal hybrid pipeline. Rather than relying solely on visual geometry "
    "to differentiate between structurally similar containers, which is difficult when physical silhouettes overlap, the framework "
    "uses semantic text extraction to resolve ambiguity without requiring model re-training."
)
set_para_clean_text(doc.paragraphs[269], p269_text)

# -------------------------------------------------------------
# Chapter 6: Conclusion
# -------------------------------------------------------------
# P275
p275_text = (
    "The rapid growth of digital media has placed significant strain on traditional forensic examination workflows. Forensic "
    "analysts are routinely tasked with identifying relevant or illicit materials within multi-terabyte storage devices containing "
    "hundreds of thousands of files. This dissertation examined the operational constraints of exact cryptographic hash matching "
    "(which is vulnerable to minor file modifications) and deep learning architectures (which require specialized computing hardware "
    "rarely available across all frontline units)."
)
set_para_clean_text(doc.paragraphs[275], p275_text)

# P276: Remove em-dash
p276_text = (
    "The central hypothesis of this research proposed that deterministic classical computer vision descriptors, specifically the "
    "Histogram of Oriented Gradients (HOG) and 3D HSV colour histograms, can deliver practical forensic triage accuracy while operating "
    "at significantly lower CPU inference latency than deep neural networks."
)
set_para_clean_text(doc.paragraphs[276], p276_text)

# P277
p277_text = (
    "The empirical findings of this study support this hypothesis. Evaluated on a curated dataset comprising Drugs, Alcohol, and Weapons "
    "categories, the Support Vector Machine (SVM) pipeline achieved an overall classification accuracy of 88.4%. More importantly, the "
    "model achieved a recall rate of 86.1%, reducing false negatives on critical evidentiary material."
)
set_para_clean_text(doc.paragraphs[277], p277_text)

# P278
p278_text = (
    "When benchmarked against transfer-learned Convolutional Neural Networks, models such as ResNet-50 attained higher raw accuracy "
    "(99.50% compared to 88.4%). However, CPU latency benchmarks demonstrated a clear throughput advantage: the classical SVM processed "
    "a single image in approximately 96.52 ms on a standard Intel CPU, whereas ResNet-50 required 294.22 ms (a threefold increase) "
    "and VGG-16 required 512.40 ms (over fivefold). When extrapolated to an investigative triage workload of 500,000 images on standard "
    "field hardware, the classical SVM pipeline completes processing in approximately 13.4 hours (achievable within an overnight operation), "
    "compared to 40.8 hours for ResNet-50 and 71.2 hours for VGG-16. This latency disparity demonstrates that classical pipelines remain "
    "viable and advantageous for frontline triage where workstation resources and time are constrained."
)
set_para_clean_text(doc.paragraphs[278], p278_text)

# P279
p279_text = (
    "Furthermore, this research demonstrated the effective integration of an Optical Character Recognition (OCR) verification module. "
    "By incorporating OCR text extraction from physical packaging, the hybrid pipeline overcomes visual ambiguities in purely geometric "
    "classifiers. Combining the computational speed of classical descriptors with semantic text parsing produces an explainable, "
    "verifiable screening tool that addresses key interpretability concerns associated with deep learning models in forensic contexts."
)
set_para_clean_text(doc.paragraphs[279], p279_text)

# P281
p281_text = (
    "Beyond technical performance metrics, rapid triage pipelines offer practical benefits for judicial workflows. In the United Kingdom "
    "and internationally, digital forensic units face substantial casework backlogs, with submitted devices often awaiting examination "
    "for extended periods. These delays can prolong pre-trial custody and delay investigative resolutions. By reducing triage times "
    "from several days to under 14 hours on standard CPU hardware, police constabularies can expedite initial evidence reviews without "
    "requiring capital investment in dedicated GPU infrastructure."
)
set_para_clean_text(doc.paragraphs[281], p281_text)

# P282
p282_text = (
    "Furthermore, automated preliminary screening supports examiner well-being. Manually reviewing large volumes of seized digital "
    "media is cognitively demanding and can expose examiners to vicarious trauma when investigating distressing cases. Using an automated "
    "classifier to filter out clearly benign background files allows analysts to focus attention on relevant evidentiary subsets, "
    "reducing prolonged manual exposure to sensitive media."
)
set_para_clean_text(doc.paragraphs[282], p282_text)

# P286
p286_text = (
    "The low computational overhead of the HOG-SVM pipeline also provides opportunities for edge deployment. Future iterations of this "
    "work could investigate porting the OpenCV pipeline to mobile platforms used by frontline officers. Optimizing descriptor computation "
    "for mobile processors could allow on-scene preliminary screening of seized devices, reducing the immediate need to transport devices "
    "to centralized laboratories for initial screening. This could accelerate preliminary intelligence assessment during on-scene inquiries."
)
set_para_clean_text(doc.paragraphs[286], p286_text)

# P295: Has math tags! Replace em-dash in run 1
p295 = doc.paragraphs[295]
for r in p295.runs:
    r.bold = False
p295.runs[1].text = p295.runs[1].text.replace(
    "—roughly five times the lifetime emissions of an average American car.",
    " (approximately five times the estimated lifetime emissions of an average passenger vehicle)."
)

# P300: Replace 'bulletproof'
p300_text = (
    "Because blockchain relies on sequential cryptographic verification, it becomes mathematically impossible for an unauthorized party "
    "to retroactively alter a flagged image or modify the classification log without invalidating the cryptographic chain. By pairing "
    "deterministic computer vision with tamper-evident logging, law enforcement agencies can provide a robust, auditable chain of custody "
    "demonstrating both the reproducibility of the triage algorithm and the integrity of the underlying digital evidence."
)
set_para_clean_text(doc.paragraphs[300], p300_text)

# P302
p302_text = (
    "Finally, future research should bridge technical computing with forensic policy and procedural law. As automated triage systems "
    "assist human investigators, standard operational procedures must govern the admissibility of machine-assisted triage findings. "
    "Further legal and technical studies should establish formal validation criteria under the Daubert and Frye standards, ensuring that "
    "laboratory performance translates into transparent, verifiable casework in court."
)
set_para_clean_text(doc.paragraphs[302], p302_text)

# P304
p304_text = (
    "In conclusion, the transition toward automated forensic evidence triage offers significant operational advantages while requiring "
    "careful attention to legal reliability and algorithmic accountability. As digital evidence volumes continue to grow, forensic "
    "frameworks must balance processing speed with evidential integrity. By combining deterministic mathematical descriptors, "
    "targeted semantic text extraction, and clear audit protocols, automated triage can assist law enforcement in managing extensive "
    "digital media while upholding procedural transparency and evidentiary standards in court."
)
set_para_clean_text(doc.paragraphs[304], p304_text)

# Save document
output_path = 'Pereira_Alston_24850898.docx'
doc.save(output_path)
print(f"Successfully saved updated dissertation to {output_path}")
