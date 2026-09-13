import docx, re

doc = docx.Document('Pereira_Alston_24850898.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")

ai_buzzwords = [
    'delve', 'delves', 'testament', 'beacon', 'pivotal', 'landscape', 
    'tapestry', 'crucial', 'staggering', 'escalating data crisis', 
    'blazingly', 'hyper-critical', 'infinitely', 'bulletproof', 
    'profound', 'drastically', 'foster', 'underscore', 'underscores',
    'vital role', 'paramount', 'holistic', 'culmination', 'realm'
]

for i, p in enumerate(doc.paragraphs):
    txt = p.text
    # check dashes
    dashes = [m.start() for m in re.finditer(r'[—–]', txt)]
    # check bold runs
    bolds = [r.text for r in p.runs if r.bold and not p.style.name.startswith('Heading')]
    # check buzzwords
    found_buzz = [w for w in ai_buzzwords if re.search(r'\b' + re.escape(w) + r'\b', txt, re.I)]
    
    if dashes or bolds or found_buzz:
        print(f"P{i:3d} (len {len(txt)}): dashes={len(dashes)}, bolds={len(bolds)}, buzzwords={found_buzz}")
        if bolds:
            print(f"     Bolds: {bolds}")
        if dashes:
            print(f"     Snippet near dash: {[txt[max(0, d-20):min(len(txt), d+25)] for d in dashes]}")
